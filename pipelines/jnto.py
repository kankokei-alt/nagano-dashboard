"""JNTO（日本政府観光局）「訪日外客統計」の国籍/月別 訪日外客数を取り込む。

JNTO の統計ページにある「国籍/月別 訪日外客数（2003年～）」の Excel を読む。
ファイル名は更新のたびに変わるので、ページからリンクを探して取得する。
シートは年ごと（2003〜最新年）。行=国・地域（「アジア計」などの地域計と、「中東地域」の内訳を含む）、
列=1〜12月と伸率・累計。

値の種類（各シートの注記と書式で判定）:
  確定 … 「すべて確定値」と注記された年
  暫定 … 最新年のうち斜体でない値
  推計 … 斜体の値（直近の月。地域計は出ないので空欄）

公表値との照合（合わない年は使わずに理由を記録）:
  1) 各年の「総数」の各月の合計 = シートの「累計」（年計）列（推計値は百人単位の丸めがあるので ±50人）
  2) 確定年の「総数」の累計 = JNTO「年別 訪日外客数、出国日本人数の推移」PDF の年の値

出力:
  data/processed/jnto_monthly.parquet … ym, country, parent, kind(total/region/country/sub), value, status
  data/processed/checks/jnto.csv       … 年ごとの照合結果

使い方:
  python pipelines/jnto.py
  python pipelines/jnto.py --file x.xlsx --annual y.pdf   # 手元のファイルを使う
"""
import argparse
import re
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
import pandas as pd
import pdfplumber
import requests

ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://www.jnto.go.jp/statistics/data/visitors-statistics/"
RAW = ROOT / "data/raw/jnto"
OUT = ROOT / "data/processed/jnto_monthly.parquet"
CHECKS = ROOT / "data/processed/checks/jnto.csv"
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) nagano-dashboard data pipeline"}


def _link(html: str, text: str, ext: str) -> str:
    m = re.search(rf'href="([^"]+\.{ext})"[^>]*>[^<]*{text}', html)
    if not m:
        raise RuntimeError(f"JNTO のページに「{text}」の {ext} が見つかりません。ページ構成が変わった可能性があります。")
    return urljoin(PAGE, m.group(1))


def download() -> tuple[Path, Path]:
    html = requests.get(PAGE, headers=UA, timeout=60).content.decode("utf-8", "replace")
    RAW.mkdir(parents=True, exist_ok=True)
    out = []
    for text, ext, name in [("国籍/月別 訪日外客数", "xlsx", "monthly.xlsx"),
                            ("年別　訪日外客数", "pdf", "annual.pdf")]:
        url = _link(html, text, ext)
        r = requests.get(url, headers=UA, timeout=180)
        r.raise_for_status()
        (RAW / name).write_bytes(r.content)
        print(f"downloaded {url}")
        out.append(RAW / name)
    return out[0], out[1]


def annual_totals(pdf: Path) -> dict[int, int]:
    """年別推移 PDF から {年: 訪日外客数}。行は「2019 平成31 / 令和元 年 31,882,049 2.2 …」。"""
    vals = {}
    with pdfplumber.open(pdf) as p:
        for line in "\n".join(pg.extract_text() or "" for pg in p.pages).splitlines():
            m = re.match(r"^(19|20)(\d\d)\s.*?年\s+([\d,]+)\s", line)
            if m:
                vals[int(m.group(1) + m.group(2))] = int(m.group(3).replace(",", ""))
    return vals


def parse_sheet(ws, year: int) -> tuple[pd.DataFrame, dict]:
    rows = list(ws.iter_rows())
    head = next(i for i, r in enumerate(rows) if any(c.value == "1月" for c in r))
    cols, total_col = {}, None
    for c in rows[head]:
        m = re.fullmatch(r"(\d+)月", str(c.value or ""))
        if m:
            cols[c.column - 1] = int(m.group(1))
        elif c.value in ("累計", "年計"):
            total_col = c.column - 1
    notes = " ".join(str(r[0].value) for r in rows if r and isinstance(r[0].value, str) and r[0].value.startswith("注"))
    final = "確定値" in notes and "推計値" not in notes

    recs, sheet_total, region = [], None, None
    for r in rows[head + 1:]:
        a, b = r[0].value, r[1].value if len(r) > 1 else None
        if isinstance(a, str) and a.startswith("注"):
            break
        # 新しい年は「中東地域」の内訳（イスラエルなど）を2列目に字下げして載せている
        if not isinstance(a, str) and isinstance(b, str) and b.strip():
            name, kind, parent = b.strip(), "sub", region
        elif isinstance(a, str) and a.strip():
            name, parent = a.strip(), None
            kind = "total" if name == "総数" else "region" if name.endswith("計") else "country"
            if kind == "country":
                region = name
        else:
            continue
        for j, mth in cols.items():
            c = r[j]
            if isinstance(c.value, (int, float)):
                st = "確定" if final else "推計" if c.font.i else "暫定"
                recs.append((pd.Timestamp(year, mth, 1), name, parent, kind, float(c.value), st))
        if kind == "total" and total_col is not None:
            sheet_total = r[total_col].value
    df = pd.DataFrame(recs, columns=["ym", "country", "parent", "kind", "value", "status"])
    return df, {"final": final, "sheet_total": sheet_total}


def build(xlsx: Path, pdf: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    official = annual_totals(pdf)
    wb = openpyxl.load_workbook(xlsx)  # 斜体（推計値）を見るので read_only にしない
    parts, checks = [], []
    for name in wb.sheetnames:
        if not name.isdigit():
            continue
        year = int(name)
        df, meta = parse_sheet(wb[name], year)
        tot = df[df.kind == "total"]
        months = len(tot)
        msum = tot.value.sum()
        # 推計値は百人単位に丸めて公表されるので、推計を含む年は累計との差 50人までを一致とみなす
        tol = 50 if (tot.status == "推計").any() else 0.5
        ok_sum = meta["sheet_total"] is not None and abs(msum - meta["sheet_total"]) <= tol
        ok_pdf = None
        if meta["final"] and months == 12:
            ok_pdf = official.get(year) == round(msum)
        ok = ok_sum and ok_pdf is not False
        reason = "" if ok else ("月の合計がシートの累計と合わない" if not ok_sum
                                else f"年別推移PDFの値（{official.get(year)}）と合わない")
        checks.append({"year": year, "months": months, "monthly_sum": msum, "sheet_total": meta["sheet_total"],
                       "annual_pdf": official.get(year), "status": "確定" if meta["final"] else "暫定・推計",
                       "adopted": ok, "reason": reason})
        print(f"  {year}: {months}か月 / 合計 {msum:,.0f} / 累計 {meta['sheet_total']} / PDF {official.get(year)}"
              f" → {'採用' if ok else '除外: ' + reason}")
        if ok:
            parts.append(df)
    data = pd.concat(parts, ignore_index=True).sort_values(["ym", "kind", "country"]).reset_index(drop=True)
    return data, pd.DataFrame(checks).sort_values("year")


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--file", type=Path, help="取得済みの国籍/月別 Excel")
    p.add_argument("--annual", type=Path, help="取得済みの年別推移 PDF")
    a = p.parse_args()
    xlsx, pdf = (a.file, a.annual) if a.file and a.annual else download()
    df, checks = build(xlsx, pdf)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    CHECKS.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(OUT, index=False)
    checks.to_csv(CHECKS, index=False)
    tot = df[df.kind == "total"]
    print(f"wrote {OUT} ({len(df):,} rows) / 総数 {tot.ym.min():%Y-%m}〜{tot.ym.max():%Y-%m}"
          f"（うち推計 {(tot.status == '推計').sum()}か月）")


if __name__ == "__main__":
    main()
